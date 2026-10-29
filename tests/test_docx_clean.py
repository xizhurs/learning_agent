import zipfile

import docx

from rag.ingest.docx_clean import accept_tracked_changes

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

BODY = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document {W}><w:body>
<w:p><w:r><w:t>Keep. </w:t></w:r>
  <w:ins w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:t>INSERTED text.</w:t></w:r></w:ins>
  <w:del w:id="2" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:delText> DELETED text.</w:delText></w:r></w:del>
  <w:r><w:t> Tail.</w:t></w:r></w:p>
<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid>
  <w:tr><w:tc><w:p><w:r><w:t>kept row</w:t></w:r></w:p></w:tc></w:tr>
  <w:tr><w:trPr><w:del w:id="3" w:author="a" w:date="2026-01-01T00:00:00Z"/></w:trPr>
    <w:tc><w:p><w:del w:id="4" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>dead row</w:delText></w:r></w:del></w:p></w:tc></w:tr>
</w:tbl>
</w:body></w:document>"""

CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

RELS = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""


def _make_docx(path):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", CONTENT_TYPES)
        z.writestr("_rels/.rels", RELS)
        z.writestr("word/document.xml", BODY)


def test_accepts_insertions_drops_deletions(tmp_path):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    _make_docx(src)
    stats = accept_tracked_changes(src, dst)

    d = docx.Document(dst)
    text = d.paragraphs[0].text
    assert "INSERTED text." in text
    assert "DELETED" not in text
    assert text.startswith("Keep.") and text.endswith("Tail.")

    xml = zipfile.ZipFile(dst).read("word/document.xml").decode()
    assert "w:ins" not in xml and "w:del" not in xml and "delText" not in xml

    assert len(d.tables[0].rows) == 1  # deleted row removed
    assert stats.insertions_accepted == 1 and stats.rows_removed == 1
