import importlib.util
import json
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from rag.config import Settings

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "export_eval_report.py"
spec = importlib.util.spec_from_file_location("export_eval_report", SCRIPT)
exp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(exp)

NOT_FOUND = "I could not find this in the provided documents."
LONG = "x" * 3000
SHEETS = [
    "summary_by_run",
    "per_question_results",
    "retrieval_details",
    "answer_details",
    "failure_analysis",
    "config",
    "paired_comparison",
    "difficulty_check",
]


def chunk(doc, n=0, text="some passage text"):
    return {"id": f"{doc}::{n:05d}", "citation": f"{doc} | n/a", "score": 1.0, "text": text}


def row(id_, answerable=True, gold=("d1.txt",), retrieved=(), answer=None, rank=None, question="q?", ragas=None, **kw):
    r = {
        "id": id_,
        "question": question,
        "type": "factual",
        "answerable": answerable,
        "source": gold[0] if gold else "",
        "reference_answer": "ref answer" if answerable else "",
        "reference_quote": "ref answer",
        "retrieved": list(retrieved),
        "rank": rank,
        "answer": answer,
        "seconds": 2.0,
        "timings": {"total": 2.0},
        "trace_id": "ab" * 16,
        "span_id": "cd" * 8,
    }
    if gold is not None:
        r["gold_context_ids"] = list(gold)
    if ragas:
        r["ragas"] = ragas
    return r | kw


def write_run(folder: Path, rows, mode="fts", size=512, top_k=5, rerank=False, testset="mytest"):
    folder.mkdir(parents=True)
    cfg = Settings().model_dump()
    cfg["retrieval"].update(mode=mode, top_k=top_k)
    cfg["chunking"].update(max_tokens=size, overlap=64)
    cfg["rerank"].update(enabled=rerank)
    (folder / "config.json").write_text(json.dumps(cfg), encoding="utf-8")
    (folder / "summary.json").write_text(json.dumps({"testset": testset, "variant": folder.name}), encoding="utf-8")
    (folder / "items.jsonl").write_text("\n".join(json.dumps(r) for r in rows), encoding="utf-8")


@pytest.fixture
def runs(tmp_path):
    root = tmp_path / "runs"
    ab = root / "20261011-100000_ablation_mytest"
    scored = [
        row(
            "q1",
            retrieved=[chunk("d1.txt"), chunk("d2.txt")],
            rank=1,
            answer="Set X [1].",
            ragas={"faithfulness": 0.9, "factual_correctness(mode=f1)": 0.8},
        ),
        row("q2", gold=("d3.txt",), retrieved=[chunk("d1.txt"), chunk("d2.txt")], answer=NOT_FOUND),
        row(
            "q3", retrieved=[chunk("d2.txt"), chunk("d2.txt", 1), chunk("d4.txt"), chunk("d1.txt", 2)], answer=NOT_FOUND
        ),
        row("q4", answerable=False, gold=(), retrieved=[chunk("d2.txt")], answer="Use option X."),
        row("q5", answerable=False, gold=(), retrieved=[chunk("d2.txt")], answer=NOT_FOUND),
        row(
            "q6",
            retrieved=[chunk("d1.txt", text=LONG)],
            answer=LONG,
            question="bell\x07 char " + LONG,
            rank=1,
            ragas={"faithfulness": 0.2, "factual_correctness(mode=f1)": 0.9},
        ),
    ]
    write_run(ab / "heading_aware_512_fts_rr-off", scored)
    unscored = [row(r["id"], r["answerable"], tuple(r.get("gold_context_ids", ())), r["retrieved"]) for r in scored]
    write_run(ab / "heading_aware_256_vector_rr-on", unscored, mode="vector", size=256, rerank=True)
    return root


def test_discovery_and_only_filter(runs):
    assert [exp.run_id_of(d, runs) for d in exp.discover_runs(runs)] == [
        "20261011-100000_ablation_mytest/heading_aware_256_vector_rr-on",
        "20261011-100000_ablation_mytest/heading_aware_512_fts_rr-off",
    ]
    assert len(exp.discover_runs(runs, "*512_fts*")) == 1 and exp.discover_runs(runs, "*nothing*") == []


def test_sheets_columns_and_summary_values(runs):
    frames, raw = exp.build_report(runs, commit="abc1234")
    assert list(frames) == SHEETS
    assert list(frames["summary_by_run"].columns) == exp.SUMMARY_COLS
    assert list(frames["per_question_results"].columns) == exp.PER_QUESTION_COLS
    assert list(frames["retrieval_details"].columns) == exp.RETRIEVAL_COLS
    assert list(frames["answer_details"].columns) == exp.ANSWER_COLS
    assert list(frames["failure_analysis"].columns) == exp.FAILURE_COLS
    assert list(frames["config"].columns) == exp.CONFIG_COLS
    assert len(raw) == 12

    s = frames["summary_by_run"].set_index("run_id")
    a = s.loc["20261011-100000_ablation_mytest/heading_aware_512_fts_rr-off"]
    assert (
        a["num_questions"] == 6 and a["retriever_type"] == "fts" and a["chunk_size"] == 512 and a["reranker"] == "none"
    )
    # answerable: q1 rank 1, q2 miss, q3 rank 4, q6 rank 1  -> recall@1 0.5, @3 0.5, @5 0.75; mrr (1 + 0 + .25 + 1) / 4
    assert (a["context_recall_at_1"], a["context_recall_at_3"], a["context_recall_at_5"]) == (0.5, 0.5, 0.75)
    assert a["mrr"] == pytest.approx(0.5625) and a["recall_basis"] == "document"
    assert a["not_found_accuracy"] == 0.5  # q5 refused, q4 answered
    assert a["ragas_faithfulness"] == pytest.approx(0.55)  # q1 0.9, q6 0.2
    assert a["answer_accuracy"] == pytest.approx(0.5)  # q1 correct, q6 correct; q2 and q3 refused (wrong); -> 2/4
    assert a["avg_latency_sec"] == 2.0 and a["estimated_cost"] == 0.0
    b = s.loc["20261011-100000_ablation_mytest/heading_aware_256_vector_rr-on"]
    assert (
        b["reranker"] == "BAAI/bge-reranker-base"
        and b["answer_accuracy"] is None
        or b["answer_accuracy"] != b["answer_accuracy"]
    )
    cfg = frames["config"].set_index("run_id").loc["20261011-100000_ablation_mytest/heading_aware_512_fts_rr-off"]
    assert (
        cfg["git_commit"] == "abc1234"
        and cfg["created_at"] == "2026-10-11T10:00:00"
        and cfg["llm_model"] == "llama3.2:3b"
    )


def test_failure_types_and_per_question_fields(runs):
    frames, _ = exp.build_report(runs, commit="")
    pq = frames["per_question_results"]
    scored = pq[pq["run_id"].str.endswith("512_fts_rr-off")].set_index("question_id")
    assert scored.loc["q1", "failure_type"] == "ok" and scored.loc["q1", "citation_correct"] is True
    assert scored.loc["q2", "failure_type"] == "retrieval_miss"
    assert scored.loc["q3", "failure_type"] == "false_refusal" and "late_hit" in scored.loc["q3", "notes"]
    assert scored.loc["q3", "gold_context_rank"] == 4 and scored.loc["q3", "gold_context_hit_at_3"] is False
    assert scored.loc["q4", "failure_type"] == "missed_refusal" and scored.loc["q4", "gold_context_hit_at_1"] is None
    assert scored.loc["q5", "failure_type"] == "ok"
    assert scored.loc["q6", "failure_type"] == "low_faithfulness"
    assert scored.loc["q1", "retrieved_context_ids"] == ["d1.txt", "d2.txt"]

    unscored = pq[pq["run_id"].str.endswith("256_vector_rr-on")].set_index("question_id")
    assert unscored.loc["q1", "failure_type"] == "ok" and unscored.loc["q2", "failure_type"] == "retrieval_miss"
    assert (
        unscored.loc["q3", "failure_type"] == "late_hit" and unscored.loc["q4", "failure_type"] == "ok"
    )  # retrieval-only: no answers

    fa = frames["failure_analysis"]
    assert set(fa["failure_type"]) == {
        "retrieval_miss",
        "false_refusal",
        "missed_refusal",
        "low_faithfulness",
        "late_hit",
    }
    assert (fa["diagnosis"] != "").all() and (fa["suggested_fix"] != "").all()
    assert len(frames["retrieval_details"]) == 2 * (2 + 2 + 4 + 1 + 1 + 1)


def test_written_files_truncate_text_freeze_header_and_keep_raw(runs, tmp_path):
    frames, raw = exp.build_report(runs, commit="")
    out = tmp_path / "reports"
    paths = exp.write_report(frames, raw, out)
    assert {p.name for p in paths} == {
        "rag_ablation_report.xlsx",
        "per_question_results.csv",
        "failure_analysis.csv",
        "raw_results.jsonl",
    }

    wb = openpyxl.load_workbook(out / "rag_ablation_report.xlsx")
    assert wb.sheetnames == SHEETS
    for ws in wb.worksheets:
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref == ws.dimensions
        for r in ws.iter_rows(values_only=True):
            assert all(len(str(v)) <= 1000 for v in r if v is not None)
    assert wb["per_question_results"].max_row == 1 + 12 and wb["summary_by_run"].max_row == 3
    assert "\x07" not in "".join(str(c.value) for c in wb["per_question_results"]["D"])  # control character removed

    csv_lines = (out / "per_question_results.csv").read_text(encoding="utf-8-sig").splitlines()
    assert len(csv_lines) >= 13 and (out / "failure_analysis.csv").exists()
    records = [json.loads(l) for l in (out / "raw_results.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(records) == 12
    q6 = next(r for r in records if r["row"]["id"] == "q6" and r["row"]["answer"])
    assert len(q6["row"]["answer"]) == 3000 and "config" in q6 and q6["derived"]["failure_type"] == "low_faithfulness"


def test_gold_ids_fall_back_to_source_and_stored_rank(tmp_path):
    root = tmp_path / "runs"
    old = [row("a", gold=None, retrieved=[chunk("d9.txt")], rank=2, source="d9.txt")]
    write_run(root / "run1", old)
    frames, _ = exp.build_report(root, commit="")
    pq = frames["per_question_results"].iloc[0]
    assert pq["gold_context_ids"] == ["d9.txt"] and pq["gold_context_rank"] == 1  # source used as the gold document
    old2 = [row("b", gold=None, source="", retrieved=[chunk("d9.txt")], rank=3)]
    write_run(root / "run2", old2)
    frames, _ = exp.build_report(root, only="run2", commit="")
    assert frames["per_question_results"].iloc[0]["gold_context_rank"] == 3
    assert frames["summary_by_run"].iloc[0]["recall_basis"] == "answer_quote"


def test_empty_runs_dir_exits_with_a_hint(tmp_path):
    with pytest.raises(SystemExit):
        exp.build_report(tmp_path, commit="")


def test_wilson_interval_and_sign_test():
    low, high = exp.wilson(45, 48)
    assert low == pytest.approx(0.8316, abs=1e-3) and high == pytest.approx(
        0.9784, abs=1e-3
    )  # 1 question is ~2 points, the CI ~15
    assert exp.wilson(48, 48)[1] == 1.0 and exp.wilson(0, 0) == (None, None)
    assert exp.sign_test_p(0, 0) == 1.0 and exp.sign_test_p(3, 3) == 1.0
    assert exp.sign_test_p(5, 0) == pytest.approx(0.0625) and exp.sign_test_p(0, 5) == pytest.approx(0.0625)
    assert exp.sign_test_p(8, 2) == pytest.approx(0.1094, abs=1e-4)


def test_paired_comparison_against_the_best_run_of_each_testset():
    runs = [
        {
            "run_id": "a",
            "testset": "t",
            "mrr": 0.9,
            "hits": {"q1": (True, True), "q2": (True, True), "q3": (False, True)},
        },
        {
            "run_id": "b",
            "testset": "t",
            "mrr": 0.7,
            "hits": {"q1": (False, True), "q2": (True, True), "q3": (True, False), "q4": (True, True)},
        },
        {"run_id": "c", "testset": "other", "mrr": 0.5, "hits": {"q1": (True, True)}},
    ]
    rows = exp.paired_comparison(runs)
    assert [(r["run_id"], r["reference_run"]) for r in rows] == [
        ("b", "a")
    ]  # 'other' has one run, so nothing to compare
    b = rows[0]
    assert b["n_paired"] == 3  # q4 is not in the reference run
    assert (b["hit1_wins"], b["hit1_losses"]) == (1, 1) and b["hit1_sign_test_p"] == 1.0
    assert (b["hit5_wins"], b["hit5_losses"]) == (0, 1)


def test_summary_has_confidence_interval_and_answer_text_recall(runs):
    rows = [
        row("a", retrieved=[chunk("d1.txt")], rank=1, answer_rank=1),
        row("b", retrieved=[chunk("d9.txt"), chunk("d1.txt")], answer_rank=None),
        row("c", retrieved=[chunk("d9.txt"), chunk("d1.txt")], answer_rank=2),
        row("d", answerable=False, gold=(), retrieved=[chunk("d1.txt")]),
    ]
    write_run(runs / "extra", rows)
    frames, _ = exp.build_report(runs, only="extra", commit="")
    s = frames["summary_by_run"].iloc[0]
    assert s["n_answerable"] == 3 and s["context_recall_at_1"] == pytest.approx(0.3333)
    assert (s["recall_at_1_ci95_low"], s["recall_at_1_ci95_high"]) == exp.wilson(1, 3)
    assert (s["answer_text_recall_at_1"], s["answer_text_recall_at_3"]) == (
        pytest.approx(0.3333),
        pytest.approx(0.6667),
    )
    pq = frames["per_question_results"].set_index("question_id")
    assert pq.loc["c", "answer_text_rank"] == 2 and pd.isna(pq.loc["d", "answer_text_rank"])
    old, _ = exp.build_report(runs, only="*512_fts*", commit="")  # saved before answer_rank existed
    assert pd.isna(old["summary_by_run"].iloc[0]["answer_text_recall_at_5"])


def test_title_baseline_ranks_documents_by_words_shared_with_the_question(tmp_path):
    folder = tmp_path / "tiny"
    folder.mkdir()
    docs = [
        {"context_id": "g1.txt", "role": "gold", "text": "Title: Oracle JDBC connectionProperties datasource\n\nbody"},
        {"context_id": "g2.txt", "role": "gold", "text": "Title: Logout button disabled in Content Navigator\n\nbody"},
        {"context_id": "x1.txt", "role": "distractor", "text": "Title: Install Datacap on Windows\n\nbody"},
        {"context_id": "x2.txt", "role": "distractor", "text": "Title: Oracle upgrade\n\nbody"},
    ]
    (folder / "contexts.jsonl").write_text("\n".join(json.dumps(d) for d in docs), encoding="utf-8")
    questions = {
        "q1": {
            "question": "How to set Oracle JDBC connectionProperties in a datasource?",
            "gold_context_ids": ["g1.txt"],
        },
        "q2": {"question": "Is the Logout button disabled?", "gold_context_ids": ["g2.txt"]},
        "q3": {"question": "Unrelated words entirely", "gold_context_ids": ["x1.txt"], "is_impossible": False},
        "q4": {"question": "Oracle", "gold_context_ids": [], "is_impossible": True},
    }
    b = exp.title_baseline("tiny/dev", questions, folder / "contexts.jsonl")
    assert b["n_answerable"] == 3 and b["n_documents"] == 4 and b["gold_share_of_documents"] == 0.5
    assert b["context_recall_at_1"] == pytest.approx(0.6667) and b["context_recall_at_5"] == 1.0
    assert exp.title_baseline("tiny/dev", questions, folder / "missing.jsonl") is None


def test_citation_check():
    retrieved = [chunk("d1.txt"), chunk("d2.txt")]
    assert exp.citation_check("A [1].", retrieved, ["d1.txt"]) is True
    assert exp.citation_check("A [2].", retrieved, ["d1.txt"]) is False  # cites a non-gold chunk
    assert exp.citation_check("A [5].", retrieved, ["d1.txt"]) is False  # out of range
    assert exp.citation_check("A without citation", retrieved, ["d1.txt"]) is False
    assert (
        exp.citation_check(NOT_FOUND, retrieved, ["d1.txt"]) is None and exp.citation_check(None, retrieved, []) is None
    )


def test_summary_has_qasper_metrics_and_leaves_them_empty_for_old_runs(runs):
    quote = "alpha beta gamma delta epsilon"
    hit = chunk("d1.txt", text="intro then alpha beta gamma delta epsilon and more")
    miss = chunk("d9.txt", text="nothing relevant here at all")
    rows = [
        row(
            "a",
            retrieved=[miss, hit],
            rank=1,
            answer_rank=2,
            answer="Set X [1].",
            reference_answer="set x",
            reference_quote=quote,
            timings={"retrieve": 1.0, "rerank": 0.8, "generate": 3.0},
            ragas={
                "answer_relevancy": 0.7,
                "context_recall": 1.0,
                "llm_context_precision_with_reference": 0.5,
                "factual_correctness(mode=f1)": None,
            },
        ),
        row(
            "b",
            retrieved=[miss, hit],
            rank=1,
            answer_rank=None,
            answer=NOT_FOUND,
            reference_quote=quote,
            timings={"retrieve": 3.0, "rerank": 1.2, "generate": 5.0},
        ),
        row(
            "c",
            answerable=False,
            gold=(),
            retrieved=[miss],
            answer=NOT_FOUND,
            timings={"retrieve": 2.0, "rerank": 1.0, "generate": 4.0},
        ),
        row("d", answerable=False, gold=(), retrieved=[miss], answer="Some answer.", timings={}),
    ]
    write_run(runs / "qasper", rows)
    frames, _ = exp.build_report(runs, only="qasper", commit="")
    s = frames["summary_by_run"].iloc[0]
    assert s["answer_mrr"] == pytest.approx(0.25)  # a: 1/2, b: miss -> (0.5 + 0) / 2
    assert (s["evidence_recall_at_1"], s["evidence_recall_at_3"], s["evidence_recall_at_5"]) == (
        0.0,
        1.0,
        1.0,
    )  # hit at rank 2
    assert s["context_precision"] == pytest.approx(0.5)  # relevant chunk at rank 2: precision@2 = 1/2
    assert s["answer_f1"] == pytest.approx(0.5)  # a: 1.0, b: refusal text vs "ref answer" -> 0
    assert s["refusal_rate_on_unanswerable"] == 0.5 and s["false_refusal_rate_on_answerable"] == 0.5
    assert s["ragas_answer_relevancy"] == pytest.approx(0.7) and s["ragas_n"] == 1
    assert pd.isna(s["ragas_factual_correctness"])  # a failed RAGAS job is skipped, not counted as 0
    assert s["median_rerank_sec"] == pytest.approx(1.0) and s["median_generate_sec"] == pytest.approx(4.0)
    pq = frames["per_question_results"].set_index("question_id")
    assert pq.loc["a", "answer_f1"] == pytest.approx(1.0) and pq.loc["a", "evidence_recall_at_1"] == 0.0
    assert pd.isna(pq.loc["c", "context_precision"]) and pd.isna(pq.loc["c", "answer_f1"])

    old = exp.build_report(runs, only="*256_vector*", commit="")[0]["summary_by_run"].iloc[
        0
    ]  # no answers, answer_rank or ragas
    assert all(
        pd.isna(old[c])
        for c in ("answer_mrr", "answer_f1", "ragas_n", "refusal_rate_on_unanswerable", "median_rerank_sec")
    )


def test_answer_correctness_prefers_the_ragas_critic_over_factual_correctness_f1(tmp_path):
    f1_zero = {"faithfulness": 0.9, "factual_correctness(mode=f1)": 0.0}
    rows = [
        row(
            "yes",
            answer="Yes, according to passage [1], they do.",
            reference_answer="Yes",
            retrieved=[chunk("d1.txt")],
            rank=1,
            ragas=f1_zero | {"answer_matches_reference": 1.0},
        ),  # right answer, F1 punishes the extra words
        row(
            "wrong",
            answer="No, they do not [1].",
            reference_answer="Yes",
            retrieved=[chunk("d1.txt")],
            rank=1,
            ragas=f1_zero | {"answer_matches_reference": 0.0},
        ),
        row(
            "old",
            answer="Yes, they do [1].",
            reference_answer="Yes",
            retrieved=[chunk("d1.txt")],
            rank=1,
            ragas=f1_zero,
        ),
        row(
            "refused",
            answer=NOT_FOUND,
            retrieved=[chunk("d1.txt")],
            rank=1,
            ragas=f1_zero | {"answer_matches_reference": 1.0},
        ),
    ]
    write_run(tmp_path / "runs" / "r", rows)
    frames, _ = exp.build_report(tmp_path / "runs", commit="")
    pq = frames["per_question_results"].set_index("question_id")
    assert list(pq.loc[["yes", "wrong", "old", "refused"], "answer_correct"]) == [True, False, False, False]
    assert list(pq.loc[["yes", "wrong", "old", "refused"], "correct_source"]) == [
        "ragas",
        "ragas",
        "ragas_factual_correctness",
        "refusal",
    ]
    assert pq.loc["yes", "failure_type"] == "ok" and pq.loc["wrong", "failure_type"] == "wrong_answer"
    assert pq.loc["old", "failure_type"] == "wrong_answer"  # old runs keep the f1 fallback
    fail = frames["failure_analysis"].set_index("question_id")
    assert "answer_matches_reference 0" in fail.loc["wrong", "diagnosis"]
    s = frames["summary_by_run"].iloc[0]
    assert s["ragas_answer_matches_reference"] == pytest.approx(
        0.667, abs=1e-3
    )  # yes 1, wrong 0, refused 1 (not excluded: it is scored)
    assert s["ragas_factual_correctness"] == 0.0 and s["answer_accuracy"] == 0.25


def test_claude_judge_score_stays_separate_from_ragas_columns(tmp_path):
    run = tmp_path / "runs" / "r"
    write_run(
        run,
        [
            row(
                "a",
                answer="Set X [1].",
                retrieved=[chunk("d1.txt")],
                rank=1,
                ragas={"faithfulness": 0.8, "answer_matches_reference": 0.0},
            )
        ],
    )
    (run / "judge_reply.json").write_text(json.dumps([{"id": "a", "correctness": 5}]), encoding="utf-8")
    pq = exp.build_report(tmp_path / "runs", commit="")[0]["per_question_results"].iloc[0]
    assert pq["claude_correctness_score"] == 1.0 and pq["ragas_answer_matches_reference"] == 0.0
    assert (  # judge replies outrank RAGAS
        pq["answer_correct"] == True and pq["correct_source"] == "claude"  # noqa: E712
    )
